"""scaffold_results.py: empty results folder and workbooks for 15.07, in the 15.06 layout (Scout, Oct 21, 2026).

Chad, Oct 21: "Scaffold the results folder and xlsx the same as we did in 15.06." Each workbook is built from its
15.06 twin (same sheet types, column names, header styles, widths and column definitions), with no data rows and
sheets only for the three 15.07 models. Two columns are new for this test set: "End token" and "Stop reason".
Run once; it refuses to overwrite a workbook that already exists.
"""
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SRC = HERE.parent / "15.06_OnDeviceApp" / "00_ondeviceResults"
OUT = HERE / "00_primeCandidateResults"

# Chad's order, Oct 21: Gemma 3 1B, Qwen3.5, MiniCPM-V 4.7. Number = workbook prefix.
MODELS = [("02", "Gemma3-1B"), ("03", "Qwen3.5"), ("04", "MiniCPM-V4.7")]
CASES = [("1-patient EOS", "Jonathan's visit alone, end token allowed, cap 2,000"),
         ("team EOS", "team two-visit case, end token allowed, cap 2,000")]
NEW_COLS = {"End token": "allowed = the model may stop at its own end token (ignore_eos 0); blocked = forced to the token count",
            "Stop reason": "eos = the model stopped by itself; n_predict = it hit the 2,000-token cap; error = decoding failed"}


def header_of(ws):
    return [c.value for c in ws[1] if c.value is not None]


def write_header(ws, cols, tmpl):
    """Header row styled like the template's header cell of the same name (or its first header cell)."""
    by_name = {c.value: c for c in tmpl[1] if c.value is not None}
    for i, name in enumerate(cols, 1):
        src = by_name.get(name, tmpl.cell(1, 1))
        dst = ws.cell(1, i, name)
        dst.font, dst.fill, dst.border, dst.alignment = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment)
        w = tmpl.column_dimensions[src.column_letter].width
        ws.column_dimensions[get_column_letter(i)].width = w if w else 14
    ws.freeze_panes = tmpl.freeze_panes


def with_new_cols(cols, after_load_mode=True):
    out = []
    for c in cols:
        out.append(c)
        if c == "Load mode" and after_load_mode:
            out.append("End token")
        if c == "Result":
            out.append("Stop reason")
    return out


def save(wb, path):
    if path.exists():
        raise SystemExit(f"exists, not overwritten: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print("wrote", path.relative_to(HERE), [ws.title for ws in wb])


def columns_sheet(wb, tmpl_wb, extra_rows=()):
    t = tmpl_wb["Columns"]
    ws = wb.create_sheet("Columns")
    write_header(ws, header_of(t), t)
    rows = [(r[0], r[1]) for r in t.iter_rows(min_row=2, values_only=True) if r[0]]
    meaning = dict(rows)
    order = with_new_cols([r[0] for r in rows]) + [r[0] for r in extra_rows]
    meaning.update(NEW_COLS)
    meaning.update(dict(extra_rows))
    for i, name in enumerate(order, 2):
        ws.cell(i, 1, name)
        ws.cell(i, 2, meaning.get(name, ""))


def raw_workbooks():
    t_wb = openpyxl.load_workbook(SRC / "01_rawResults" / "00_OnDeviceResults.xlsx")
    tmpl = t_wb["Qwen3.5 Q8_0 - llama.cpp CPU"]
    cols = with_new_cols(header_of(tmpl))

    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for _, m in MODELS:
        for case, _ in CASES:
            write_header(wb.create_sheet(f"{m} Q8_0 {case}"[:31]), cols, tmpl)
    write_header(wb.create_sheet("Rig checks"), cols, t_wb["Rig checks"])
    columns_sheet(wb, t_wb)
    save(wb, OUT / "01_rawResults" / "00_PrimeCandidateResults.xlsx")

    # Per-model workbooks: that model's sheets plus the full generated text, like 15.06's 02..06 _Results
    m_wb = openpyxl.load_workbook(SRC / "01_rawResults" / "02_Qwen3.5_Results.xlsx")
    m_tmpl = m_wb["Qwen3.5 Q8_0 - llama.cpp CPU"]
    m_cols = with_new_cols(header_of(m_tmpl))
    for num, m in MODELS:
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        for case, _ in CASES:
            write_header(wb.create_sheet(f"{m} Q8_0 {case}"[:31]), m_cols, m_tmpl)
        columns_sheet(wb, m_wb)
        save(wb, OUT / "01_rawResults" / f"{num}_{m}_Results.xlsx")


def quality_workbook():
    t_wb = openpyxl.load_workbook(SRC / "01_rawResults" / "01_OnePatientRuns_Quality.xlsx")
    tmpl = t_wb["Runs"]
    cols = []
    for c in header_of(tmpl):
        if c.startswith("Generated note"):
            cols += ["Stop reason", "Generated note (full text, end token allowed, cap 2,000)"]
        else:
            cols.append(c)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    write_header(wb.create_sheet("Runs"), cols, tmpl)
    p = wb.create_sheet("Prompt")
    write_header(p, header_of(t_wb["Prompt"]), t_wb["Prompt"])
    keep = {r[0]: r[1] for r in t_wb["Prompt"].iter_rows(min_row=2, values_only=True) if r[0]}
    rows = [("Test", "One-patient note quality, end token allowed (Chad, Oct 21: let the model do what the model does). NOT a speed row: tokens out differ per model."),
            ("Conversation", keep["Conversation"]),
            ("Instructions", keep["Instructions"]),
            ("Settings", "Each model's fastest CPU setting from the 15.06 speed runs; end token allowed (ignore_eos 0), cap 2,000 tokens; greedy; context 10,240; app scout-bench f23b516; two runs per model."),
            ("SYNTHETIC-ENGINEERING-ONLY_gemma3-1b_jonathan.txt", keep["SYNTHETIC-ENGINEERING-ONLY_gemma3-1b_jonathan.txt"]),
            ("SYNTHETIC-ENGINEERING-ONLY_qwen35_jonathan.txt", keep["SYNTHETIC-ENGINEERING-ONLY_qwen35_jonathan.txt"]),
            ("MiniCPM-V 4.7 1B", keep["MiniCPM-V 4.7 1B"])]
    for i, (k, v) in enumerate(rows, 2):
        p.cell(i, 1, k); p.cell(i, 2, v)
    save(wb, OUT / "01_rawResults" / "01_OnePatientRuns_Quality.xlsx")


def processing_workbooks():
    """Team-template copies (15.01 02_WindowsBenchmarks layout): a settings sheet plus a '- runs' sheet per case."""
    t_wb = openpyxl.load_workbook(SRC / "00_processingResults" / "02_Qwen3.5_Benchmarks.xlsx")
    s_tmpl, r_tmpl = t_wb["Qwen3.5 Q8_0 CPU"], t_wb["Qwen3.5 Q8_0 CPU - runs"]
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for _, m in MODELS:
        write_header(wb.create_sheet(f"{m} Q8_0 1-patient EOS"[:31]), header_of(s_tmpl), s_tmpl)
    for _, m in MODELS:
        name = f"{m} 1-patient EOS - runs"
        write_header(wb.create_sheet(name if len(name) <= 31 else f"{m} 1p EOS - runs"), header_of(r_tmpl), r_tmpl)
    save(wb, OUT / "00_processingResults" / "01_OnePatientRuns_Benchmarks.xlsx")
    for num, m in MODELS:
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        write_header(wb.create_sheet(f"{m} Q8_0 team EOS"[:31]), header_of(s_tmpl), s_tmpl)
        write_header(wb.create_sheet(f"{m} team EOS - runs"[:31]), header_of(r_tmpl), r_tmpl)
        save(wb, OUT / "00_processingResults" / f"{num}_{m}_Benchmarks.xlsx")


def summary_workbook():
    """Same sections as 15.06's summary: title, verdict, device, engine, at a glance, one block per model, every run."""
    t_wb = openpyxl.load_workbook(SRC / "00_OnDeviceBenchmarkSummary.xlsx")
    t = t_wb.active
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Prime Candidate Results"
    for k, d in t.column_dimensions.items():
        ws.column_dimensions[k].width = d.width

    def put(row, values, style_row, merge=False):
        for i, v in enumerate(values, 1):
            c = ws.cell(row, i, v); s = t.cell(style_row, i)
            c.font, c.fill, c.border, c.alignment = copy(s.font), copy(s.fill), copy(s.border), copy(s.alignment)
        if merge:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)

    put(1, ["Local LLM Benchmarks: Prime Model Candidate (Galaxy A17)"], 1)
    put(2, ["3 models (Gemma 3 1B IT, Qwen3.5-0.8B, MiniCPM-V 4.7 1B), all Q8_0 on llama.cpp. Benchmarks to show the best candidate for fine-tuning later; no fine-tuning in this set."], 2, merge=True)
    put(4, ["Verdict"], 4)
    put(5, ["Pending: no runs yet."], 5, merge=True)
    put(7, ["Device tested"], 8)
    put(8, [t.cell(9, i).value for i in range(1, 9)], 9)
    put(9, [t.cell(10, i).value for i in range(1, 9)], 10)
    put(11, ["Engine"], 12)
    put(12, ["Engine", "What it is", "Q8_0 possible?", "Runs that ran", "Outcome in one line"], 13)
    put(13, ["llama.cpp a0e6dc7", "llama.cpp's own Android example app (llama.android) with our test mode; app scout-bench f23b516 (adds the end-token switch)", "Yes", "0 so far", "Pending"], 14)
    put(15, ["At a glance: one row per model (Q8_0, llama.cpp CPU, end token allowed)"], 17)
    put(16, ["Model", "Setting", "KV cache (K/V)", "One-patient: tokens written", "One-patient: stop reason", "One-patient: time to finished note (s)",
             "Team case: tokens written", "Team case: stop reason", "Team case: time to finished note (s)", "Peak RAM (GB)", "Meets 60 s?", "Note quality (fact check)", "Notes"], 18)
    settings = {"Gemma3-1B": ("Gemma 3 1B IT", "CPU, 2 threads, flash attention on", "q8_0/q8_0"),
                "Qwen3.5": ("Qwen3.5-0.8B", "CPU, 2 threads, flash attention off", "q8_0/f16"),
                "MiniCPM-V4.7": ("MiniCPM-V 4.7 1B", "CPU, 2 threads, flash attention off", "f16/f16")}
    r = 17
    for _, m in MODELS:
        put(r, list(settings[m]) + ["pending"], 19); r += 1
    put(r, ["Times are seconds, load excluded. Tokens written differ per model, so these are not speed rows; the 15.06 speed rows (1,000 tokens forced) stay the speed comparison."], 24, merge=True)
    r += 2
    put(r, ["By model"], 26); r += 1
    for i, (_, m) in enumerate(MODELS, 1):
        put(r, [f"{i}. {settings[m][0]}"], 27); r += 1
        put(r, [t.cell(28, c).value for c in range(1, 14)], 28); r += 1
        for case, desc in CASES:
            put(r, [f"{case.replace(' EOS', '')}, llama.cpp, end token allowed", settings[m][1], "Q8_0", settings[m][2], "0 of 2", "", "", "", "", "", "", "", desc], 29); r += 1
        r += 1
    put(r, ["Every run"], 26); r += 1
    put(r, ["Run ID", "Model", "Case", "Setting", "Stop reason", "Tokens written", "Load time (s)", "Time to first token (s)",
            "Prompt reading (tok/s)", "Writing speed (tok/s)", "Total time, load excluded (s)", "Peak RAM (GB)", "Result"], 28)
    save(wb, OUT / "00_PrimeCandidateSummary.xlsx")


def readme():
    p = OUT / "README.txt"
    if p.exists():
        raise SystemExit(f"exists, not overwritten: {p}")
    p.write_text("""PRIME MODEL CANDIDATE RESULTS (Samsung Galaxy A17)
Benchmarks to show which of 3 models is the best candidate for fine-tuning later. No fine-tuning here.
Models (Chad, Oct 21, 2026): 02 Gemma3-1B, 03 Qwen3.5, 04 MiniCPM-V4.7. All Q8_0 on llama.cpp, CPU.
Set up Oct 21, 2026 in the same layout as 15.06_OnDeviceApp\\00_ondeviceResults. No runs yet.

FOLDER LAYOUT
  00_PrimeCandidateSummary.xlsx      START HERE. Verdict, one row per model, each model's block, every run.
  charts\\                            Charts of the same results (added once there are runs).
  00_processingResults\\              Results in the team's benchmark template.
  01_rawResults\\                     Every measurement, plus the full text each model wrote.

00_processingResults\\
  01_OnePatientRuns_Benchmarks.xlsx  One-patient note test, in the template.
  02 to 04 _<Model>_Benchmarks.xlsx  One file per model (team case). Each setting averaged over its two runs,
                                     plus a "- runs" sheet with each run.

01_rawResults\\
  00_PrimeCandidateResults.xlsx      Every run, every column. One sheet per model + case.
  01_OnePatientRuns_Quality.xlsx     One-patient note test. Full note text per run.
  02 to 04 _<Model>_Results.xlsx     One file per model: its rows from 00 plus the full generated text.

WHAT THE SHEET NAMES MEAN
  1-patient EOS   Jonathan's visit alone. The model may stop at its own end token (cap 2,000 tokens).
  team EOS        The team two-visit case, same end-token rule.
  End token / Stop reason columns are new in this set: did the model stop by itself, or hit the cap?
  These are NOT speed rows (each model writes a different number of tokens).
  Times are seconds, load time excluded. Target is 60 s.

Per-run raw files (output text, logs): ..\\results\\
""", encoding="utf-8")
    print("wrote", p.relative_to(HERE))


if __name__ == "__main__":
    (OUT / "charts").mkdir(parents=True, exist_ok=True)
    raw_workbooks()
    quality_workbook()
    processing_workbooks()
    summary_workbook()
    readme()
